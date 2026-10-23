import io

import pytest
from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

from app.services.cv_parser import CVParseError, parse_cv

LINES = [
    "Jane Doe - Software Engineer",
    "Experience: Led migration of billing service to Kubernetes, cutting costs by 30%.",
    "Skills: Python, FastAPI, PostgreSQL, React",
]


def make_pdf(lines) -> bytes:
    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in lines:
        c.drawString(50, y, line)
        y -= 20
    c.save()
    return buf.getvalue()


def make_docx(lines, table_rows=()) -> bytes:
    doc = Document()
    for line in lines:
        doc.add_paragraph(line)
    if table_rows:
        t = doc.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, val in enumerate(row):
                t.cell(r, c).text = val
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def test_pdf():
    text = parse_cv(make_pdf(LINES), "cv.PDF")
    assert "Kubernetes" in text and "Jane Doe" in text


def test_docx_with_table_in_order():
    text = parse_cv(make_docx(LINES, [("2020-2023", "Acme Corp, Senior Engineer")]), "cv.docx")
    assert "billing service" in text
    assert "2020-2023 | Acme Corp, Senior Engineer" in text
    assert text.index("Skills") < text.index("Acme Corp")


def test_unsupported_type():
    with pytest.raises(CVParseError, match="PDF or DOCX"):
        parse_cv(b"hello", "cv.txt")


def test_empty_pdf_hints_at_scan():
    with pytest.raises(CVParseError, match="scanned"):
        parse_cv(make_pdf([]), "cv.pdf")


def test_corrupt_file():
    with pytest.raises(CVParseError, match="valid"):
        parse_cv(b"not a real pdf", "cv.pdf")


def test_unmapped_pdf_glyphs_become_bullets():
    from app.services.cv_parser import normalise

    assert normalise("(cid:127) Led migration\n(cid:127)Built API (cid:3)") == "• Led migration\n• Built API"
